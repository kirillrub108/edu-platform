"""Bulk roster import for a restricted course.

Parses an .xlsx/.csv of student emails and hands every address to
`course_access_service.grant_access` — the same write path the single
"add by email" form uses, so there is exactly one way a grant is created.
The file is parsed in memory and discarded; nothing is stored.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from redis.asyncio import Redis
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.constants import ROSTER_IMPORT_MAX_ROWS, ROSTER_IMPORT_SAMPLE_VALUES
from app.models.course import Course, CourseAccessGrant
from app.models.user import User, UserRole
from app.schemas.course import (
    CourseAccessGrantCreate,
    RosterColumnOption,
    RosterColumnRef,
    RosterImportReport,
    RosterImportRow,
)
from app.services import course_access_service


class RosterImportError(ValueError):
    """Unusable upload — the router turns this into a 400 with its message."""


# Anything that is clearly not a spreadsheet/CSV. cp1251 decodes almost every
# byte sequence, so a magic check is what actually stops a .pdf renamed to .csv.
_BINARY_MAGICS: tuple[bytes, ...] = (
    b"%PDF-",
    b"\xd0\xcf\x11\xe0",  # legacy .xls / .doc (OLE2)
    b"\x89PNG",
    b"\xff\xd8\xff",  # JPEG
    b"\x1f\x8b",  # gzip
)
_ZIP_MAGIC = b"PK\x03\x04"

_EMAIL_HEADERS: frozenset[str] = frozenset(
    {
        "email",
        "e-mail",
        "mail",
        "почта",
        "эл. почта",
        "эл почта",
        "электронная почта",
        "адрес электронной почты",
    }
)

# NFKC folds NBSP and full-width forms; zero-width characters survive it and
# have to be dropped explicitly (they ride along in copy-pasted address lists).
_ZERO_WIDTH = dict.fromkeys([0x200B, 0x200C, 0x200D, 0x2060, 0xFEFF], None)


# ── Parsing ─────────────────────────────────────────────────────────────────


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return value if isinstance(value, str) else str(value)


def _too_many_rows(count: int) -> RosterImportError:
    return RosterImportError(
        f"В файле {count} строк — максимум {ROSTER_IMPORT_MAX_ROWS}. "
        "Разделите список на несколько файлов."
    )


def _collect(rows: Iterable[list[str]]) -> list[list[str]]:
    """Keep at most ROSTER_IMPORT_MAX_ROWS rows but keep *counting* past the
    cap, so the error can name the real row count without holding the file."""
    grid: list[list[str]] = []
    count = 0
    for row in rows:
        count += 1
        if count <= ROSTER_IMPORT_MAX_ROWS:
            grid.append(row)
    if count > ROSTER_IMPORT_MAX_ROWS:
        raise _too_many_rows(count)
    return grid


def _parse_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zipfile/KeyError/ValueError variants
        raise RosterImportError("Не удалось прочитать XLSX — файл повреждён") from exc
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise RosterImportError("В книге нет ни одного листа")
        return _collect([_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def _decode_csv(content: bytes) -> str:
    # utf-8-sig first (Excel's own UTF-8 export), cp1251 last — Russian Excel's
    # default "CSV (разделители — запятые)" encoding.
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise RosterImportError(
        "Не удалось определить кодировку CSV (ожидается UTF-8 или Windows-1251)"
    )


def _parse_csv(content: bytes) -> list[list[str]]:
    if b"\x00" in content[:4096]:
        raise RosterImportError("Это не CSV — файл содержит двоичные данные")
    text = _decode_csv(content)
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return _collect(list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter))


def parse_grid(filename: str, content: bytes) -> list[list[str]]:
    """Sync parser (openpyxl / csv). Format comes from the extension AND the
    content signature — never from the browser-supplied content type."""
    if not content.strip():
        raise RosterImportError("Файл пустой")
    name = filename.lower()
    is_zip = content.startswith(_ZIP_MAGIC)
    if name.endswith(".xlsx"):
        if not is_zip:
            raise RosterImportError("Файл не похож на XLSX")
        return _parse_xlsx(content)
    if name.endswith(".csv"):
        if is_zip or content.startswith(_BINARY_MAGICS):
            raise RosterImportError("Файл не похож на CSV")
        return _parse_csv(content)
    raise RosterImportError("Поддерживаются только файлы .xlsx и .csv")


# ── Column detection ────────────────────────────────────────────────────────


def _header_key(raw: str) -> str:
    folded = unicodedata.normalize("NFKC", raw).lower().replace("ё", "е")
    return " ".join(folded.split()).strip(":*")


def normalize_email(raw: str) -> str:
    value = unicodedata.normalize("NFKC", raw).translate(_ZERO_WIDTH)
    value = " ".join(value.split())
    if "<" in value and ">" in value:  # "Имя Фамилия <mail@example.com>"
        value = value[value.index("<") + 1 : value.rindex(">")]
    value = value.strip().strip("\"'")
    if value.lower().startswith("mailto:"):
        value = value[len("mailto:") :]
    return value.strip().lower()


def validate_email(raw: str) -> str | None:
    """Same validator as the single add-by-email form, via its request schema."""
    candidate = normalize_email(raw)
    if not candidate:
        return None
    try:
        return CourseAccessGrantCreate(email=candidate).email.lower()
    except ValidationError:
        return None


def _column_options(
    grid: list[list[str]], header_row: list[str] | None
) -> list[RosterColumnOption]:
    width = max((len(row) for row in grid), default=0)
    options: list[RosterColumnOption] = []
    body = grid[1:] if header_row is not None else grid
    for index in range(width):
        values = [row[index].strip() for row in body if index < len(row) and row[index].strip()]
        titled = bool(header_row and index < len(header_row) and header_row[index].strip())
        if not values and not titled:
            continue
        header = ""
        if header_row is not None and index < len(header_row):
            header = header_row[index].strip()
        options.append(
            RosterColumnOption(
                index=index,
                header=header or get_column_letter(index + 1),
                samples=values[:ROSTER_IMPORT_SAMPLE_VALUES],
            )
        )
    return options


@dataclass(frozen=True)
class _Selection:
    index: int | None
    header: str
    skip_first_row: bool
    options: list[RosterColumnOption]


def select_column(grid: list[list[str]], explicit: int | None) -> _Selection:
    """1) a header cell that reads like an email column; 2) failing that, the
    single column that actually holds valid addresses; 3) otherwise ask."""
    first = grid[0] if grid else []
    header_hit = next(
        (i for i, cell in enumerate(first) if _header_key(cell) in _EMAIL_HEADERS), None
    )
    if explicit is not None:
        if explicit < 0 or explicit >= max((len(row) for row in grid), default=0):
            raise RosterImportError("Указанной колонки нет в файле")
        index = explicit
    elif header_hit is not None:
        index = header_hit
    else:
        with_emails = [
            i
            for i in range(max((len(row) for row in grid), default=0))
            if any(i < len(row) and validate_email(row[i]) for row in grid)
        ]
        if len(with_emails) != 1:
            # For the choice UI a first row holding no address at all is shown
            # as column titles, even if none of them says "email".
            titles = first if len(grid) > 1 and not any(validate_email(c) for c in first) else None
            return _Selection(None, "", False, _column_options(grid, titles))
        index = with_emails[0]

    # A header row is only skipped when it is really a header: its own cell is
    # not an address while some later row's is. Without that, a file whose
    # first line is already data would lose its first student.
    body_has_email = any(index < len(row) and validate_email(row[index]) for row in grid[1:])
    first_is_email = index < len(first) and validate_email(first[index]) is not None
    skip_first = bool(grid) and not first_is_email and body_has_email
    header = first[index].strip() if skip_first and index < len(first) else ""
    return _Selection(index, header or get_column_letter(index + 1), skip_first, [])


# ── Import ──────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class _Candidate:
    row: int
    raw: str
    email: str | None


async def import_from_file(
    db: AsyncSession,
    redis: Redis,
    course: Course,
    teacher: User,
    filename: str,
    content: bytes,
    email_column: int | None,
) -> RosterImportReport:
    grid = await run_in_threadpool(parse_grid, filename, content)
    selection = select_column(grid, email_column)
    if selection.index is None:
        return RosterImportReport(needs_column_choice=True, columns=selection.options)

    column = selection.index
    body = grid[1:] if selection.skip_first_row else grid
    offset = 2 if selection.skip_first_row else 1  # report rows as seen in Excel
    candidates = [
        _Candidate(row=n + offset, raw=row[column].strip(), email=validate_email(row[column]))
        for n, row in enumerate(body)
        if column < len(row) and row[column].strip()
    ]

    wanted = {c.email for c in candidates if c.email}
    students: dict[str, User] = {}
    if wanted:
        found = await db.scalars(
            select(User).where(
                func.lower(User.email).in_(wanted),
                User.role == UserRole.student,
            )
        )
        students = {user.email.lower(): user for user in found}
    granted: set[object] = set()
    if students:
        rows = await db.scalars(
            select(CourseAccessGrant.student_id).where(
                CourseAccessGrant.course_id == course.id,
                CourseAccessGrant.student_id.in_([u.id for u in students.values()]),
            )
        )
        granted = set(rows)

    owner_email = teacher.email.lower()
    seen: set[str] = set()
    results: list[RosterImportRow] = []

    def skip(candidate: _Candidate, reason: str) -> None:
        results.append(
            RosterImportRow(
                row=candidate.row,
                raw_value=candidate.raw,
                email=candidate.email,
                status="skipped",
                reason=reason,  # type: ignore[arg-type]
            )
        )

    for candidate in candidates:
        if candidate.email is None:
            skip(candidate, "invalid_email")
            continue
        if candidate.email in seen:
            skip(candidate, "duplicate_in_file")
            continue
        seen.add(candidate.email)
        if candidate.email == owner_email:
            skip(candidate, "is_course_owner")
            continue
        student = students.get(candidate.email)
        if student is None:
            skip(candidate, "not_registered")
            continue
        if student.id in granted:
            skip(candidate, "already_in_list")
            continue
        # ponytail: one commit per student (grant_access commits) — the single
        # add-by-email path reused verbatim. Batch it only if 1000-row imports
        # become slow in practice.
        await course_access_service.grant_access(db, course.id, student.id, teacher.id)
        await course_access_service.publish_access_change(redis, student.id, course.id, "granted")
        granted.add(student.id)
        results.append(
            RosterImportRow(
                row=candidate.row,
                raw_value=candidate.raw,
                email=candidate.email,
                status="added",
            )
        )

    added = sum(1 for r in results if r.status == "added")
    return RosterImportReport(
        detected_column=RosterColumnRef(index=column, header=selection.header),
        total_rows=len(results),
        added=added,
        skipped=len(results) - added,
        results=results,
    )
